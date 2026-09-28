import json,collections
from pathlib import Path
from openpyxl import load_workbook
p=Path('NewData/data/processed/company_knowledge.jsonl'); rows=[json.loads(s) for s in p.read_text(encoding='utf8').splitlines()]
print('NEW CATEGORIES',collections.Counter(r['metadata']['category'] for r in rows))
for cat in dict.fromkeys(r['metadata']['category'] for r in rows):
    r=next(r for r in rows if r['metadata']['category']==cat);print(cat,json.dumps(r,ensure_ascii=False)[:2700])
k=json.loads(Path('.local-internal/knowledge.json').read_text(encoding='utf8'))
print('INTERNAL',list(k),collections.Counter(d['kind'] for d in k['documents']))
for d in k['documents']:
    if d['kind']=='service':print('SERVICE',d['title'],d['body'][:350])
for p in Path('NewData/data/workbooks').glob('*.xlsx'):
    w=load_workbook(p,read_only=True,data_only=False);print('WORKBOOK',p.name,[(s.title,s.max_row,s.max_column) for s in w]);w.close()
for p in ['data/demo_documents.json','data/workflow_documents.json','data/security_documents.json']:
    ds=json.loads(Path(p).read_text(encoding='utf8'));print('OLD',p,collections.Counter(d['category'] for d in ds)); print('SAFE IDS',[(d['id'],d['title']) for d in ds if not d.get('customer')][:160])
