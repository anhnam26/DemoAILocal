"""Read-only checks for a synchronized, theory-only working dataset."""
import json,sqlite3
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent
docs=json.loads((ROOT/'data/knowledge_documents.json').read_text(encoding='utf8'))
forbidden=('CRM-','CONTRACT-','PROJECT-','QUOTE-','TICKET-','COST-','AR-','INV-','PAY-','CASE-')
assert len({d['id'] for d in docs})==len(docs)
assert not any(d.get('customer') or d.get('is_example') or d['id'].startswith(forbidden) for d in docs)
assert not (ROOT/'.local-internal').exists()
assert not (ROOT/'NewData/data/indexes').exists()
assert not (ROOT/'NewData/data/legacy').exists()
assert not list((ROOT/'NewData/data/processed').glob('example*'))
for path in (ROOT/'NewData/data/workbooks').glob('*.xlsx'):
    wb=load_workbook(path,read_only=True)
    assert not any(s.startswith('example_') or s=='Du_lieu_MINH_HOA' for s in wb.sheetnames)
    wb.close()
with sqlite3.connect(ROOT/'data/demo.sqlite3') as c:
    stored=[json.loads(r[0]) for r in c.execute('SELECT payload FROM docs')]
    assert not any(d.get('customer') or d['id'].startswith(forbidden) for d in stored)
    assert not c.execute("SELECT 1 FROM users WHERE role NOT IN ('member','admin') OR customers!='[]'").fetchone()
    assert c.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    assert {d['id'] for d in docs} <= {d['id'] for d in stored}
print(json.dumps({'corpus_documents':len(docs),'database_documents':len(stored),'customer_record_ids':0,'example_sheets':0,'sqlite_integrity':'ok'},indent=2))
